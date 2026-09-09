"""Tạo hướng dẫn cải thiện chi tiết cho dự án VietSign Vision."""

from pathlib import Path

from build_improvement_report import (
    BLUE,
    GRAY,
    LIGHT_BLUE,
    PALE_BLUE,
    RED,
    add_bullet,
    add_callout,
    add_table,
    configure_sections,
    configure_styles,
    set_font,
)
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "VietSign_Vision_Detailed_Improvement_Guide.docx"


def add_code(document: Document, code: str) -> None:
    """Thêm khối code ngắn, dễ sao chép."""
    table = document.add_table(rows=1, cols=1)
    from build_improvement_report import set_cell_shading, set_table_geometry

    set_table_geometry(table, [9360])
    set_cell_shading(table.cell(0, 0), "F2F4F7")
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(code.strip().splitlines()):
        run = paragraph.add_run(line)
        set_font(run, name="Consolas", size=8.5, color="1F2933")
        if index < len(code.strip().splitlines()) - 1:
            run.add_break()


def add_step(
    document: Document,
    title: str,
    reason: str,
    files: str,
    actions: list[str],
    done: list[str],
) -> None:
    document.add_heading(title, level=2)
    add_callout(document, "Vì sao", reason, fill=PALE_BLUE, accent=BLUE)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("File liên quan: ")
    run.bold = True
    paragraph.add_run(files)
    for action in actions:
        add_bullet(document, action)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Definition of Done")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("2E7D32")
    for item in done:
        add_bullet(document, f"[ ] {item}")


def build_document() -> None:
    document = Document()
    configure_styles(document)
    configure_sections(document)

    section = document.sections[0]
    section.header.paragraphs[0].clear()
    header = section.header.paragraphs[0].add_run("VIETSIGN VISION  |  DETAILED IMPROVEMENT GUIDE")
    set_font(header, size=8.5, color=GRAY, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("Hướng dẫn cải thiện chi tiết")
    set_font(run, size=27, color=BLUE, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Từ portfolio học thuật đến dự án ML có thể tái lập và bảo vệ khi phỏng vấn"
    )
    set_font(run, size=12, color=GRAY, italic=True)

    document.add_paragraph()
    add_callout(
        document,
        "Mục tiêu",
        "Sửa tính đúng đắn ML trước, sau đó mới hoàn thiện khả năng tái lập, test, tài liệu, demo và hiệu năng. Không thêm framework hoặc tính năng chỉ để dự án trông phức tạp.",
        fill=LIGHT_BLUE,
        accent=RED,
    )

    document.add_heading("1. Thứ tự ưu tiên", level=1)
    add_table(
        document,
        ["Giai đoạn", "Mức", "Kết quả cần đạt", "Ước lượng"],
        [
            ["1. Sửa split và leakage", "P0", "Metric hợp lệ, có test độc lập", "1–2 ngày"],
            ["2. Làm sạch claim", "P0", "Không có nội dung sai hoặc phóng đại", "0,5 ngày"],
            ["3. Tái lập dữ liệu/model", "P0", "Clone sạch có đường đi rõ ràng", "1–2 ngày"],
            ["4. Metric, baseline, test", "P1", "Có bằng chứng kỹ thuật đáng tin", "2–4 ngày"],
            ["5. Demo và hiệu năng", "P1", "Chứng minh use case offline", "1–2 ngày"],
            ["6. Đóng gói CV/GitHub", "P2", "README trung thực, dễ đánh giá", "1 ngày"],
        ],
        [1700, 900, 4560, 2200],
    )
    add_callout(
        document,
        "Điểm dừng hợp lý cho CV",
        "Hoàn thành Giai đoạn 1–4 đã đủ tạo một portfolio ML tốt. Chỉ làm web API/load test 100 users nếu bạn thực sự muốn ứng tuyển Backend/MLOps; với CV Engineer, benchmark offline và demo video có giá trị hơn.",
        fill="FFF7E6",
        accent="B26A00",
    )

    document.add_heading("2. Giai đoạn 1 — Sửa tính đúng đắn AI/ML", level=1)
    add_step(
        document,
        "2.1 Tạo train / validation / test độc lập",
        "Notebook hiện dùng test_files.txt làm validation rồi chọn C và threshold trên chính tập đó. Kết quả vì vậy không còn là test độc lập.",
        "data/raw/split_dataset/, notebooks/04_task4_roi.ipynb, 05_task5_hog.ipynb, 06_task6_svm.ipynb",
        [
            "Giữ test_files.txt bất biến và tuyệt đối không dùng để chọn feature, C, gamma hoặc threshold.",
            "Tách validation từ train theo ảnh nguồn. Nếu ảnh đến từ video/tuyến đường, phải group theo video/tuyến đường trước khi chia để ảnh gần nhau không nằm ở hai split.",
            "Thêm val_files.txt. Mỗi stem chỉ được xuất hiện trong đúng một split.",
            "ROI sinh từ một ảnh phải kế thừa split của ảnh đó; không chia ngẫu nhiên theo ROI.",
            "Chỉ đọc test ở bước đánh giá cuối cùng sau khi mọi tham số đã khóa.",
        ],
        [
            "train, val, test đôi một không giao nhau.",
            "Không có cell tuning nào đọc test_files.txt.",
            "Báo cáo số ảnh, ROI và phân phối lớp cho từng split.",
            "Test chỉ được chạy một lần cho bảng kết quả phát hành.",
        ],
    )
    add_code(
        document,
        """
train_stems = load_split_stems(split_dir / "train_files.txt")
val_stems = load_split_stems(split_dir / "val_files.txt")
test_stems = load_split_stems(split_dir / "test_files.txt")

assert train_stems.isdisjoint(val_stems)
assert train_stems.isdisjoint(test_stems)
assert val_stems.isdisjoint(test_stems)

split = "train" if stem in train_stems else "val" if stem in val_stems else "test"
""",
    )

    add_step(
        document,
        "2.2 Loại bỏ leakage trong GridSearchCV",
        "tune_svm() hiện fit StandardScaler trên toàn bộ X_train trước khi chia fold. Mean/std của validation fold đã ảnh hưởng dữ liệu huấn luyện trong mỗi fold.",
        "src/classifier.py, tests/test_pipeline.py",
        [
            "Đưa StandardScaler và SVC vào sklearn.pipeline.Pipeline.",
            "GridSearch tham số theo tên svc__C và svc__gamma.",
            "Sau GridSearch, lấy scaler và SVC từ best_estimator_.named_steps; pipeline đã refit trên toàn bộ train bằng tham số tốt nhất.",
            "Không gọi fit_transform bên ngoài CV.",
        ],
        [
            "GridSearchCV nhận X_train thô.",
            "Mỗi fold tự fit scaler chỉ trên training fold.",
            "Test xác nhận best estimator chứa scaler và svc.",
            "15 test cũ vẫn pass và có test regression mới.",
        ],
    )
    add_code(
        document,
        """
pipe = Pipeline([
    ("scaler", StandardScaler()),
    ("svc", SVC(kernel="rbf", class_weight="balanced",
                probability=probability, random_state=random_state)),
])
grid = {"svc__C": [0.5, 1, 2, 5, 10],
        "svc__gamma": ["scale", 0.001, 0.01]}
search = GridSearchCV(pipe, grid, scoring="f1_macro", cv=skf, n_jobs=n_jobs)
search.fit(X_train, y_train)
best = search.best_estimator_
return best.named_steps["svc"], best.named_steps["scaler"]
""",
    )

    add_step(
        document,
        "2.3 Khóa quy trình chọn threshold",
        "bin_thr và multi_thr đang được tối ưu trên tập được gọi là test. Threshold là tham số mô hình và chỉ được chọn trên validation.",
        "notebooks/06_task6_svm.ipynb, config.yaml",
        [
            "Train model trên train; chọn C/gamma bằng CV chỉ trong train.",
            "Chọn bin_thr và multi_thr trên validation theo mục tiêu rõ: ưu tiên recall hay macro F1.",
            "Khóa config, sau đó đánh giá đúng một lần trên test.",
            "Không chỉnh threshold sau khi nhìn test result.",
        ],
        [
            "Config ghi rõ nguồn của từng tham số: CV hoặc validation.",
            "Bảng test không được dùng để ra quyết định tuning.",
            "Có file JSON lưu toàn bộ sweep validation.",
        ],
    )

    document.add_heading("3. Giai đoạn 2 — Làm sạch kiến thức và claim", level=1)
    add_step(
        document,
        "3.1 Sửa attribution CLAHE",
        "Ngưỡng std 50/100 và clipLimit 4/2/1 chưa có nguồn học thuật xác minh. Giữ attribution sai làm giảm độ tin cậy toàn dự án.",
        "src/preprocessing.py, config.yaml, README.md, docs/ARCHITECTURE.md",
        [
            "Đổi mô tả thành 'heuristic nội bộ được chọn thực nghiệm'.",
            "Nếu muốn giữ citation, phải ghi DOI, tác giả, trang và chứng minh công thức trùng nguồn.",
            "Chạy ablation fixed CLAHE 1/2/4 so với heuristic trên validation và báo candidate recall + macro F1.",
            "Nếu heuristic không cải thiện ổn định, dùng clipLimit cố định đơn giản hơn.",
        ],
        [
            "Không còn chuỗi Vieira 2024 khi chưa có bibliography hợp lệ.",
            "README phân biệt rõ công thức chuẩn và heuristic của dự án.",
            "Có bảng ablation hoặc ghi rõ chưa đánh giá.",
        ],
    )

    add_step(
        document,
        "3.2 Hạ các claim vượt bằng chứng",
        "Các cụm production-grade, edge-ready, Explainable AI và Type Annotations 100% cần phép đo hoặc định nghĩa. Portfolio tốt ưu tiên trung thực hơn quảng cáo.",
        "README.md, pyproject.toml",
        [
            "Dùng 'interpretable classical CV baseline' thay cho Explainable AI nếu chưa có XAI method.",
            "Dùng 'CPU-only offline pipeline' thay cho edge-ready đến khi benchmark Raspberry Pi/phần cứng mục tiêu.",
            "Dùng 'modular, tested Python package' thay cho production-grade.",
            "Không ghi 100% type hints nếu chưa chạy mypy/pyright và đo coverage annotation.",
        ],
        [
            "Mọi con số trong README trỏ đến artifact tái lập.",
            "Không còn claim tuyệt đối không có phép đo.",
            "Phần CV nói rõ giới hạn dataset/model.",
        ],
    )

    document.add_heading("4. Giai đoạn 3 — Tái lập dữ liệu và model", level=1)
    add_step(
        document,
        "4.1 Hoàn thiện Dataset Card và manifest",
        "Repository hiện có 5/3.191 ảnh tham chiếu, không đủ train hoặc xác minh benchmark.",
        "data/raw/, src/audit.py",
        [
            "Ghi nguồn, phiên bản, license, quy trình annotation, class mapping và quyền phân phối.",
            "Nếu không thể commit data, tạo script tải hoặc hướng dẫn đặt dữ liệu với checksum.",
            "Tạo manifest CSV/JSON gồm relative path, split, group/source, SHA-256, kích thước ảnh và label count.",
            "Audit thêm trùng ảnh exact hash và near-duplicate pHash giữa các split.",
        ],
        [
            "Clone sạch có thể chuẩn bị data theo một lệnh/hướng dẫn xác định.",
            "Audit phát hiện overlap, duplicate và thiếu file.",
            "Dataset version được khóa bằng hash/manifest.",
        ],
    )

    add_step(
        document,
        "4.2 Phát hành model có provenance",
        "Full inference không chạy vì thiếu svm_binary.joblib và svm_multiclass.joblib; Joblib cũng không an toàn nếu tải từ nguồn lạ.",
        "outputs/models/, src/classifier.py, README.md",
        [
            "Phát hành model qua GitHub Release hoặc kho artifact; không commit file lớn trực tiếp.",
            "Công bố SHA-256, scikit-learn version, Python version, config hash, dataset manifest hash và training commit.",
            "Thêm lệnh verify checksum trước load.",
            "Ghi rõ chỉ load Joblib từ release chính thức; không nhận model upload tùy ý.",
        ],
        [
            "Hai model tải được từ URL ổn định.",
            "Checksum sai làm chương trình dừng với lỗi rõ.",
            "Full CLI chạy trên ảnh mẫu từ môi trường sạch.",
        ],
    )

    document.add_heading("5. Giai đoạn 4 — Metric, baseline và kiểm thử", level=1)
    document.add_heading("5.1 Bộ metric tối thiểu", level=2)
    add_table(
        document,
        ["Tầng", "Metric chính", "Metric bổ sung", "Lý do"],
        [
            [
                "Candidate",
                "Recall@IoU 0.3/0.5",
                "boxes/image, FP/image",
                "Không bỏ sót biển trước classifier",
            ],
            [
                "Binary",
                "Macro F1, sign recall",
                "PR-AUC, confusion matrix",
                "Nền thường lệch lớp; FN truyền xuống tầng 2",
            ],
            ["Multiclass", "Macro F1", "per-class P/R/F1", "52 lớp có phân phối không đều"],
            [
                "End-to-end",
                "mAP50 hoặc F1 detection",
                "accuracy chỉ bổ sung",
                "Phải tính cả vị trí box và lớp",
            ],
            ["Confidence", "ECE/Brier", "reliability diagram", "Confidence SVC cần calibration"],
            ["Runtime", "latency p50/p95", "RAM, throughput", "Chứng minh claim CPU/edge"],
        ],
        [1500, 2300, 2600, 2960],
    )

    document.add_heading("5.2 Baseline nên có", level=2)
    for item in [
        "Majority/random baseline cho classification.",
        "HOG + Linear SVM để chứng minh giá trị của RBF.",
        "HOG + RBF SVM một tầng để chứng minh giá trị kiến trúc hai tầng.",
        "Candidate không NMS so với NMS để chứng minh trade-off recall/box count.",
        "Một detector nhẹ như YOLO nano chỉ để so sánh accuracy/latency, không bắt buộc thay kiến trúc chính.",
    ]:
        add_bullet(document, item)

    add_step(
        document,
        "5.3 Mở rộng test đúng rủi ro",
        "15 test hiện tại tốt cho utility nhưng chưa bảo vệ quy trình ML và full CLI.",
        "tests/test_core.py, tests/test_pipeline.py, tests/test_audit.py, thêm tests/test_cli.py",
        [
            "Test malformed YOLO label: thiếu cột, NaN, class ngoài phạm vi, tọa độ âm/vượt 1.",
            "Test corrupt image, thư mục rỗng, classes file ngắn hơn predicted class.",
            "Test Pipeline trong GridSearch để ngăn scaler leakage quay lại.",
            "Parity test: cùng một ROI tạo HOG giống nhau ở đường training và inference.",
            "CLI integration test với detect-only và JSON schema; lỗi một file không làm mất toàn bộ batch nếu chính sách cho phép.",
            "Model compatibility test cho metadata/version/checksum.",
        ],
        [
            "Tối thiểu 22–25 test có ý nghĩa, không chạy chậm quá mức.",
            "Có một golden smoke test đầu-cuối không cần model.",
            "Có một full inference test dùng model fixture nhỏ hoặc release model.",
        ],
    )

    document.add_heading("6. Giai đoạn 5 — Demo, hiệu năng và production", level=1)
    add_step(
        document,
        "6.1 Tạo demo chứng minh use case",
        "Ảnh detect-only chỉ chứng minh candidate generation, chưa chứng minh nhận dạng 52 lớp.",
        "docs/assets/, README.md, tools/benchmark.py (mới)",
        [
            "Chọn 10–20 ảnh test chưa dùng tuning, gồm ngày/đêm, biển nhỏ, mờ, che khuất và nền phức tạp.",
            "Tạo GIF/video before-after có box, nhãn, confidence và thời gian xử lý.",
            "Thêm failure gallery: false positive, false negative và nhầm lớp; giải thích nguyên nhân.",
            "Không chỉ chọn ảnh đẹp nhất; ghi rõ cách chọn mẫu.",
        ],
        [
            "Demo full recognition chạy được bằng lệnh trong README.",
            "Có ít nhất ba failure case trung thực.",
            "Ảnh demo không thuộc train/val.",
        ],
    )

    add_step(
        document,
        "6.2 Benchmark trước khi nói edge-ready hoặc 100 users",
        "Hiện chưa có server, load test hoặc benchmark phần cứng. Không cần xây service nếu không thuộc mục tiêu dự án.",
        "tools/benchmark.py, README.md",
        [
            "Warm-up 5 lần, đo ít nhất 30 lần; báo p50/p95 thay vì một lần đo.",
            "Đo riêng candidate generation, HOG, classifier và tổng pipeline.",
            "Ghi CPU, RAM, OS, Python/OpenCV versions và kích thước ảnh.",
            "Nếu làm API: giới hạn kích thước upload, timeout, queue/backpressure, process pool và load test 1/10/100 concurrent requests.",
            "Nếu chỉ làm portfolio CV: ghi rõ offline single-process; không đưa claim 100 users.",
        ],
        [
            "Có bảng p50/p95/RAM/throughput tái lập.",
            "Claim hiệu năng luôn kèm phần cứng và input size.",
            "Không dùng từ production-ready nếu chưa có service observability và load test.",
        ],
    )

    document.add_heading("7. Giai đoạn 6 — README và CV", level=1)
    document.add_heading("7.1 Cấu trúc README đề xuất", level=2)
    for item in [
        "1. Bài toán và phạm vi: offline ảnh tĩnh, 52 lớp, mục tiêu giáo dục/baseline.",
        "2. Demo thật: ảnh/GIF và lệnh tái lập.",
        "3. Kiến trúc: sơ đồ 6 tầng và input/output schema.",
        "4. Kết quả: validation và test tách riêng, metric + hardware rõ ràng.",
        "5. Cài đặt: clone → install → download/verify model → run.",
        "6. Data/Model provenance: nguồn, license, checksum.",
        "7. Testing/CI: đúng số test hiện tại, lệnh chạy.",
        "8. Limitations & responsible use: đặt ngay sau kết quả.",
        "9. Roadmap và tài liệu chuyên sâu.",
    ]:
        add_bullet(document, item)

    document.add_heading("7.2 Bullet CV chỉ dùng sau khi tái lập", level=2)
    add_callout(
        document,
        "Mẫu tiếng Việt",
        "Xây dựng pipeline nhận dạng biển báo Việt Nam trên CPU, kết hợp candidate generation bằng HSV/MSER/Canny, HOG 1.764 chiều và SVM hai tầng; đóng gói CLI, dataset audit, kiểm thử tự động và benchmark tái lập trên test split khóa.",
        fill=PALE_BLUE,
        accent=BLUE,
    )
    add_callout(
        document,
        "Quy tắc ghi metric",
        "Chỉ thêm số Macro F1/mAP/latency khi có script, dataset manifest, model checksum và test report tái lập. Nếu chưa có, mô tả năng lực thiết kế thay vì ghi số cũ trong config.",
        fill="FFF1F0",
        accent=RED,
    )

    document.add_heading("8. Kế hoạch thực hiện 7 ngày", level=1)
    add_table(
        document,
        ["Ngày", "Công việc", "Artifact đầu ra"],
        [
            [
                "1",
                "Tạo split độc lập; sửa notebook không dùng test để tuning",
                "val_files.txt + split audit",
            ],
            ["2", "Sửa scaler leakage; thêm regression test", "classifier.py + test"],
            ["3", "Sửa claim CLAHE/README/LICENSE", "README và docs trung thực"],
            ["4", "Dataset manifest + model metadata/checksum", "manifest + Model Card"],
            ["5", "Baseline và metric pipeline", "evaluation JSON/CSV"],
            ["6", "Edge cases, CLI integration, parity test", "22–25 test pass"],
            ["7", "Demo, benchmark, README và CV bullet", "release candidate"],
        ],
        [1000, 4960, 3400],
    )

    document.add_heading("9. Lệnh kiểm tra cuối", level=1)
    add_code(
        document,
        """
python -m pip install -e ".[dev,notebooks]"
ruff check src tests tools
python -m compileall -q src tests tools
pytest -q
python -m src.audit --output outputs/dataset-audit.json
python -m src.cli data/raw/images/0589.jpg --detect-only
python -m src.cli data/raw/images/0589.jpg
python tools/benchmark.py --input data/raw/images --runs 30
""",
    )

    document.add_heading("10. Definition of Done toàn dự án", level=1)
    for item in [
        "[ ] Không dùng test để chọn feature, hyperparameter hoặc threshold.",
        "[ ] StandardScaler nằm trong Pipeline khi cross-validation.",
        "[ ] Metric test tái lập từ dataset/model có hash.",
        "[ ] Có baseline và metric detection phù hợp.",
        "[ ] Training–inference parity được kiểm thử.",
        "[ ] Clone sạch chạy detect-only và full recognition theo README.",
        "[ ] Model có checksum/provenance; LICENSE tồn tại.",
        "[ ] README không còn số test/metric/claim sai.",
        "[ ] Demo có cả thành công và failure cases.",
        "[ ] Claim latency/edge/production luôn có benchmark tương ứng.",
    ]:
        add_bullet(document, item)

    document.add_paragraph(
        "Tài liệu này dựa trên audit mã nguồn và artifact hiện có ngày 01/09/2026. Thứ tự ưu tiên cố ý đặt tính đúng đắn ML và khả năng tái lập trước giao diện hoặc mở rộng kiến trúc."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
